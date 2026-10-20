"""Read hospital workbook as data; never execute spreadsheet content."""
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sheet = load_workbook(ROOT / '核心制度监测指标统计提示词.xlsx', read_only=True, data_only=True).worksheets[0]
units = ['住院人次','住院人次','病例','住院人次','急会诊次数','急会诊次数','普通会诊次数','普通会诊次数','手术患者','四级手术患者','再次住院/手术对象','讨论对象','异常费用患者','抢救例次','手术例次','已讨论手术例次','手术例次','手术例次','死亡病例','死亡病例（不同集合）','死亡病例','应上传死亡病案','长期医嘱条目','手术例次','手术例次','发生率之比','患者死亡率之比','四级手术例次','术种','项目种类','分钟','危急值条目','特殊级抗菌药物医嘱','输血例次','术中输血患者']
rows = []
for group, indexes in [('core', range(2, 37)), ('quality', range(39, 55))]:
    for n, row in enumerate(indexes, 1):
        kind = 'rate'
        required = ['numerator', 'denominator']
        if group == 'core':
            if n == 20:
                kind = 'ratio'
            elif n in (26, 27):
                kind, required = 'rate_ratio', ['four_success', 'four_total', 'three_success', 'three_total']
            elif n == 29:
                kind, required = 'catalog_rate', ['performed_types', 'registered_types']
            elif n == 30:
                kind, required = 'retention', list('ABCDEF')
            elif n == 31:
                kind, required = 'median', ['minutes', 'eligible_total']
        notes = str(sheet.cell(row, 7).value or '')
        if group == 'core' and n == 20:
            notes += '\n已纠正：分母是发生医疗纠纷的死亡病例数，不能使用全部死亡病例数。医务部门参与不等于组织。'
        if group == 'core' and n == 22:
            notes += '\n必须有应上传清单、实际完整上传结果；死亡医嘱不能证明上传完成。'
        if group == 'core' and n == 19:
            notes += '\n工作日需要医院节假日/调休日历；缺日历时人工复核。'
        rows.append({'id':f'{group}_{n:02}', 'name':sheet.cell(row,2).value, 'definition':sheet.cell(row,3).value,
                     'formula':sheet.cell(row,4).value, 'notes':notes, 'kind':kind,
                     'numerator_prompt':str(sheet.cell(row,5).value or ''),
                     'denominator_prompt':str(sheet.cell(row,6).value or ''),
                     'hospital_notes':str(sheet.cell(row,8).value or ''),
                     'prompt_sources':{'numerator':f'{sheet.title}!E{row}', 'denominator':f'{sheet.title}!F{row}',
                                       'notes':f'{sheet.title}!G{row}', 'hospital_notes':f'{sheet.title}!H{row}'},
                     'unit':units[n-1] if group=='core' else '病例', 'required_measures':required,
                     'source':f'{sheet.title}!B{row}:G{row}',
                     'automation':'candidate_events' if group=='core' and n in (1,2) else 'human_assisted',
                     'requires_external':group=='core' and n in (20,22,29,30)})
(ROOT / 'medocr' / 'catalog.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding='utf-8')
print(f'Generated {len(rows)} indicators. Original workbook unchanged.')
